import json
from copy import deepcopy
from pathlib import Path
import sys
import unittest
import tempfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'skills/jh-math-exam/scripts'))
from exam_validation import validate_exam, ExamValidationError, source_question_digest
from generate_exam_docx import (exam_header, build_exam_paper, build_student_answer_paper,
                                build_answer_paper, Figures, validate_question_review_sources,
                                build_question_review)
from exam_validation import exam_digest
from math_omml import add_rich_text, build_omath, plain_text, M_NS


def planned_exam():
    questions = [dict(number=i+1,section='3-1',bloom_level='理解',points=4,
                      question=f'化簡 {{{i+2}*x-x}}。',answer=f'{{{i+1}*x}}',
                      solution=f'合併得 {{{i+1}*x}}。',primary_topic='T',
                      provenance=dict(kind='material',file='課本.pdf',locator='PDF第3頁'))
                 for i in range(5)]
    questions[-1]['provenance'] = dict(kind='web',url='https://example.org/question',
                                     title='公開題目',locator='例1',approval={'status':'pending'})
    return dict(grade='七年級',scope='3-1',total_points=20,fill_questions=questions,
                content_plan=[dict(id='T',section='3-1',name='合併同類項',target_points=20,rationale='核心能力')],
                source_policy=dict(basis='points',material_percent=80,web_percent=20))


class PlanningAndLayoutTests(unittest.TestCase):
    def review_sources(self, data):
        return dict(schema_version=1,exam_sha256=exam_digest(data),questions={
            f'fill_{q["number"]}':dict(relationship='concept_reference' if q['provenance']['kind']=='web' else 'concept_original',sources=[dict(
                kind='web' if q['provenance']['kind']=='web' else 'workbook',
                title='來源名稱',locator='印刷第55頁；PDF第2頁；第3題',
                original='原題標記：化簡 {3*x-x}。',transcription='原題轉錄',
                url='https://example.org/question')]) for q in data['fill_questions']})

    def test_question_review_rejects_missing_stale_and_incomplete_sources(self):
        data=planned_exam()
        for mutation in (
            lambda s: s['questions'].pop('fill_1'),
            lambda s: s.update(exam_sha256='stale'),
            lambda s: s['questions']['fill_1']['sources'][0].pop('original'),
            lambda s: s['questions']['fill_1']['sources'][0].update(kind='web'),
        ):
            sources=self.review_sources(data); mutation(sources)
            with self.assertRaises(ExamValidationError):
                validate_question_review_sources(data,sources)

    def test_question_review_has_originals_and_classifications_without_solutions(self):
        data=planned_exam(); sources=self.review_sources(data)
        for q in data['fill_questions']: q['solution']='不可印出的解析'
        sources['exam_sha256']=exam_digest(data)
        with tempfile.TemporaryDirectory() as folder:
            doc=Document(build_question_review(data,Path(folder),Figures(None),sources))
            text='\n'.join(p.text for p in doc.paragraphs)
            for part in ('第一部分：本次題目','第二部分：來源與原題','第三部分：重點分類'):
                self.assertEqual(text.count(part),5)
            self.assertIn('非原先直接改編題',text)
            self.assertIn('原題標記',text)
            self.assertIn('合併同類項',text)
            self.assertNotIn('不可印出的解析',text)
            self.assertIn('oMath',doc.element.xml)
            self.assertNotIn('{3*x-x}',text)
            self.assertEqual(doc.sections[0]._sectPr.find(qn('w:cols')).get(qn('w:num')),'2')

    def test_pending_web_blocks_formal_but_allows_review(self):
        data = planned_exam()
        self.assertTrue(validate_exam(data,allow_pending_web=True)['warnings'])
        with self.assertRaisesRegex(ExamValidationError,'禁止正式匯出'):
            validate_exam(data)
        approval = dict(status='approved',question_sha256=source_question_digest(data,'fill_5'),
                        user_confirmation='使用者明確同意候選第五題')
        data['fill_questions'][-1]['provenance']['approval']=approval
        self.assertFalse(validate_exam(data)['warnings'])
        data['fill_questions'][-1]['solution'] += '補充一步。'
        with self.assertRaises(ExamValidationError):
            validate_exam(data)

    def test_topic_coverage_and_source_proportions_cannot_be_faked_by_omission(self):
        for mutate in (
            lambda d: d['fill_questions'][0].update(primary_topic='missing'),
            lambda d: d['content_plan'][0].update(target_points=16),
            lambda d: d['fill_questions'][0]['provenance'].update(kind='web'),
            lambda d: d['fill_questions'][0].update(section='3-2'),
        ):
            data = planned_exam(); mutate(data)
            with self.assertRaises(ExamValidationError):
                validate_exam(data,allow_pending_web=True)

    def test_coefficient_and_variable_typography_preserves_numeric_multiplication(self):
        self.assertEqual(plain_text('{17*x+17*3+x*3}'),'17x+17×3+x×3')
        doc = Document(); p=doc.add_paragraph()
        add_rich_text(p,'每本x元，17×x；exam不拆字。')
        variables=[r for r in p.runs if r.text=='x']
        self.assertEqual(len(variables),2)
        self.assertTrue(all(r.italic and r.font.name=='Cambria Math' for r in variables))
        root=build_omath('17*x+3')
        variable=root.xpath('.//m:r[m:t="x"]',namespaces={'m':M_NS})[0]
        self.assertEqual(variable.find(f'{{{M_NS}}}rPr/{{{M_NS}}}sty').get(f'{{{M_NS}}}val'),'i')
        self.assertEqual(''.join(root.itertext()),'17x+3')
        self.assertNotIn('×',''.join(build_omath('17*x^2').itertext()))

    def test_material_and_web_are_aggregated_before_topic_balance_check(self):
        data=planned_exam()
        data['content_plan']=[
            dict(id='T',section='3-1',name='合併同類項',target_points=12,rationale='全卷12分'),
            dict(id='U',section='3-1',name='分配律',target_points=8,rationale='全卷8分')]
        data['fill_questions'][3]['primary_topic']='U'
        data['fill_questions'][4]['primary_topic']='U'
        self.assertEqual(validate_exam(data,allow_pending_web=True)['total_points'],20)
        # 教材仍16分、網路仍4分，但網路題改主重點後，合計16/4偏離12/8。
        data['fill_questions'][4]['primary_topic']='T'
        with self.assertRaisesRegex(ExamValidationError,'實際16分|實際 16 分'):
            validate_exam(data,allow_pending_web=True)

    def test_columns_start_after_full_width_header(self):
        doc=Document(); exam_header(doc,{'total_points':100},'題目卷')
        doc.add_paragraph('雙欄題目')
        self.assertEqual(len(doc.sections),1)
        self.assertTrue(doc.sections[0].header.tables)
        self.assertEqual(doc.sections[0]._sectPr.find(qn('w:cols')).get(qn('w:num')),'2')
        self.assertEqual(doc.sections[0]._sectPr.find(qn('w:cols')).get(qn('w:space')),'454')
        self.assertIn('＿＿＿＿國中',doc.sections[0].header.tables[0].cell(0,0).text)
        custom=Document(); exam_header(custom,{'total_points':100,'school_name':'測試國中'},'題目卷')
        self.assertIn('測試國中',custom.sections[0].header.tables[0].cell(0,0).text)
        self.assertNotIn('＿＿＿＿國中',custom.sections[0].header.tables[0].cell(0,0).text)

    def test_student_sheet_has_blank_answer_cells_and_teacher_key_excludes_solutions(self):
        from test_exam_validation import valid_exam
        data=valid_exam()
        data['fill_questions'][0]['answer']='教師專用答案'
        for questions in ('mc_questions','fill_questions'):
            for q in data[questions]: q['solution']='解析隱藏標記'
        for sub in data['open_questions'][0]['sub_questions']: sub['solution']='解析隱藏標記'
        with tempfile.TemporaryDirectory() as folder:
            student=Document(build_student_answer_paper(data,Path(folder),Figures(None)))
            for table in student.tables:
                grid=table.cell(0,0).tables[0]
                self.assertTrue(all(not cell.text.strip() for cell in grid.rows[1].cells[1:]))
            student_text=''.join(student.element.itertext())
            self.assertNotIn('教師專用答案',student_text)
            self.assertNotIn('解析隱藏標記',student_text)
            self.assertIn('每個書袋同價',student_text)
            teacher=Document(build_answer_paper(data,Path(folder)))
            teacher_text=''.join(teacher.element.itertext())
            self.assertIn('教師專用答案',teacher_text)
            self.assertNotIn('解析隱藏標記',teacher_text)
            self.assertNotIn('每個書袋同價',teacher_text)
            self.assertNotIn('常數項是何者',teacher_text)
            paper=Document(build_exam_paper(data,Path(folder),Figures(None)))
            fill=next(p for p in paper.paragraphs if '化簡' in p.text)
            self.assertNotIn('分）',fill.text)
            self.assertGreater(fill.paragraph_format.left_indent,0)
            self.assertLess(fill.paragraph_format.first_line_indent,0)
            self.assertEqual(len(paper.sections),1)


if __name__=='__main__':
    unittest.main()
