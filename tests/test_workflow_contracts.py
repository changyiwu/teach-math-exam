"""新版流程的外部契約：不能繞過來源規劃或回到舊匯出方式。"""
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from test_exam_validation import valid_exam, report_checks, SCRIPTS
from exam_validation import exam_digest, write_math_report
from generate_exam_docx import build_exam_paper, build_student_answer_paper, build_answer_paper, build_question_review, Figures


def context_exam():
    base=valid_exam()
    big=deepcopy(base['open_questions'][0])
    for sub in big['sub_questions']:
        sub.update(primary_topic='T',provenance=dict(kind='material',file='測試教材.pdf',locator='測試例1'))
    return dict(title='流程測試用情境卷',grade='七年級',scope=big['section'],
                total_points=12,scope_sections=[big['section']],expected_counts=dict(mc=0,fill=0,open=1),
                mc_questions=[],fill_questions=[],open_questions=[big],
                content_plan=[dict(id='T',name='列式與解方程式',section=big['section'],target_points=12,rationale='測試')],
                source_policy=dict(basis='points',material_percent=100,web_percent=0))


def context_sources(data):
    return dict(schema_version=1,exam_sha256=exam_digest(data),questions={
        f'open_1_{i}':dict(relationship='concept_original',sources=[dict(kind='textbook',
            title='測試教材',locator='測試例1',original='測試原題資料',transcription='僅供測試')])
        for i in (1,2)})


class WorkflowContractTests(unittest.TestCase):
    def test_context_review_resolves_shared_parent_figure_using_current_id(self):
        from PIL import Image
        data=context_exam()
        data['open_questions'][0]['geometry']=dict(spec=dict(type='triangle'))
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            Image.new('RGB',(32,32),'white').save(root/'open_1_fig.png')
            result=build_question_review(data,root,Figures(root),context_sources(data))
            self.assertEqual(len(Document(result).inline_shapes),2)

    def test_formal_export_requires_plan_sources_and_new_output_directory(self):
        data=context_exam()
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            src=root/'exam.json'
            report=root/'math.json'
            sources=root/'sources.json'
            for output,remove,expected in (
                ('flat',None,'Word/'),
                ('file.docx',None,'Word/'),
                ('Word','content_plan','主重點'),
                ('Word','review_sources','question-review-sources'),
            ):
                current=deepcopy(data)
                if remove=='content_plan':
                    current.pop('content_plan');current.pop('source_policy')
                src.write_text(json.dumps(current,ensure_ascii=False),encoding='utf-8')
                write_math_report(current,report_checks(current),report)
                sources.write_text(json.dumps(context_sources(current),ensure_ascii=False),encoding='utf-8')
                args=[sys.executable,str(SCRIPTS/'generate_exam_docx.py'),str(src),str(root/output),'--math-report',str(report)]
                if remove!='review_sources':args+=['--question-review-sources',str(sources)]
                result=subprocess.run(args,capture_output=True,text=True,encoding='utf-8')
                with self.subTest(output=output,remove=remove):
                    self.assertNotEqual(result.returncode,0)
                    self.assertIn(expected,result.stderr)
                    self.assertFalse((root/output).exists())

    def test_context_papers_are_separate_b4_documents_with_answers_only_for_teacher(self):
        data=context_exam()
        for sub in data['open_questions'][0]['sub_questions']:sub['solution']='解析不得出現在三種考卷'
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder)
            paths=[build_exam_paper(data,root,Figures(None)),
                   build_student_answer_paper(data,root,Figures(None)),build_answer_paper(data,root)]
            for path in paths:
                doc=Document(path)
                sec=doc.sections[0]
                self.assertAlmostEqual(sec.page_width.cm,25.7,places=2)
                self.assertEqual(sec._sectPr.find(qn('w:cols')).get(qn('w:num')),'2')
                self.assertIn(data['title'],sec.header.tables[0].cell(0,0).text)
                text='\n'.join(p.text for p in doc.paragraphs)
                self.assertNotIn('解析不得出現在三種考卷',text)
            teacher='\n'.join(p.text for p in Document(paths[2]).paragraphs)
            self.assertIn('6個',teacher)
            self.assertNotIn(data['open_questions'][0]['context'],teacher)
            self.assertIn('一、計算題',Document(paths[0]).paragraphs[0].text)


if __name__=='__main__':unittest.main()
